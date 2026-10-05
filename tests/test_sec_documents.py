"""Service-free: security-review fixes (docs/implementation/status.md) for documents, ingestion and storage (H5, H6, M11-M13, L9, L10, L12, L6).

No stack, secrets or network: Supabase, Neo4j and Temporal are replaced by in-memory fakes.
"""

import io
from datetime import UTC, datetime
from types import SimpleNamespace

import fitz
import openpyxl
import pytest
from fastapi import HTTPException, Response
from starlette.datastructures import Headers, UploadFile

from api.routers import documents as docs
from api.services import ocr
from api.services.graph import GraphService
from api.services.pid import PIDService
from api.services.pii import PIIService
from workflows import document_pipeline as pipeline

NOW = datetime(2026, 10, 1, tzinfo=UTC)


# ---------------------------------------------------------------------------
# In-memory fakes
# ---------------------------------------------------------------------------

class FakeQuery:
    def __init__(self, db, table):
        self.db, self.table, self.filters, self.op, self.payload = db, table, [], "select", None

    def __getattr__(self, name):
        if name in {"select", "limit", "order", "in_", "range", "not_", "neq", "gte", "lte"}:
            return lambda *a, **k: self
        raise AttributeError(name)

    def eq(self, column, value):
        self.filters.append((column, value))
        self.db.eq_calls.append((self.table, column, value))
        return self

    def insert(self, row):
        self.op, self.payload = "insert", row
        return self

    def update(self, row):
        self.op, self.payload = "update", row
        return self

    def upsert(self, row, **kwargs):
        self.op, self.payload = "upsert", row
        return self

    def execute(self):
        if self.op in ("insert", "update", "upsert"):
            self.db.writes.append((self.op, self.table, self.payload, list(self.filters)))
            row = dict(self.payload)
            if self.table == "extraction_jobs":
                row["job_id"] = "job-1"
            return SimpleNamespace(data=[row])
        rows = [r for r in self.db.tables.get(self.table, []) if all(r.get(c) == v for c, v in self.filters)]
        return SimpleNamespace(data=rows)


class FakeBucket:
    def __init__(self, db):
        self.db = db

    def upload(self, path, data, options=None):
        self.db.uploads.append((path, options))

    def remove(self, paths):
        self.db.removed.extend(paths)

    def create_signed_url(self, path, expires):
        return {"signedURL": f"https://x.supabase.co/storage/v1/object/sign/kairos-vault/{path}?token=t"}


class FakeSupabase:
    def __init__(self, **tables):
        self.tables, self.writes, self.eq_calls, self.uploads, self.removed = tables, [], [], [], []
        self.storage = SimpleNamespace(from_=lambda bucket: FakeBucket(self))

    def table(self, name):
        return FakeQuery(self, name)

    def inserted(self, table):
        return [p for op, t, p, _ in self.writes if op == "insert" and t == table]

    def updated(self, table):
        return [p for op, t, p, _ in self.writes if op == "update" and t == table]


class FakeTemporal:
    def __init__(self):
        self.started = []

    async def start_workflow(self, *args, **kwargs):
        self.started.append(kwargs["args"][0])


def upload(name="note.txt", data=b"hello", content_type="text/plain"):
    return UploadFile(io.BytesIO(data), size=len(data), filename=name,
                      headers=Headers({"content-type": content_type}))


def user(role, user_id="u-1"):
    return {"user_id": user_id, "role": role}


async def ingest(supabase, current_user, **form):
    args = {"asset_id": None, "document_type": "procedure", "source_system": "manual_upload",
            "authority_level": 4, "occurred_at": None, "file": upload()}
    args.update(form)
    temporal = FakeTemporal()
    result = await docs.ingest_document(current_user=current_user, supabase=supabase, temporal=temporal, **args)
    return result, temporal


# ---------------------------------------------------------------------------
# M12: storage path inputs
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name,expected", [
    ("../../other-bucket/secret.pdf", "secret.pdf"),
    ("..\\..\\win\\evil.txt", "evil.txt"),
    ("/abs/path/report.pdf", "report.pdf"),
    ("bulletin scan (1).png", "bulletin_scan__1_.png"),
    ("..", "upload"),
    ("", "upload"),
    (None, "upload"),
    (".hidden", "hidden"),
])
def test_filename_is_reduced_to_a_safe_leaf(name, expected):
    assert docs.safe_filename(name) == expected


async def test_ingest_builds_the_storage_path_from_vetted_parts():
    db = FakeSupabase()
    result, temporal = await ingest(db, user("admin"), file=upload("../../x/../evil name.txt"))
    path = db.uploads[0][0]
    assert path.startswith("procedure/DOC-") and path.endswith("/evil_name.txt")
    assert ".." not in path and path.count("/") == 2
    assert temporal.started[0]["vault_path"] == path == result["vault_path"]


@pytest.mark.parametrize("doc_type", ["../escape", "free text", "", "OEM_MANUAL"])
async def test_ingest_rejects_a_document_type_outside_the_fixed_list(doc_type):
    db = FakeSupabase()
    with pytest.raises(HTTPException) as exc:
        await ingest(db, user("admin"), document_type=doc_type)
    assert exc.value.status_code == 422
    assert db.uploads == []


# ---------------------------------------------------------------------------
# H5: authority, asset, occurred_at
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("role", ["engineer", "operator", "field_worker"])
async def test_a_non_designated_role_is_capped_to_level_four_and_it_is_disclosed(role):
    db = FakeSupabase()
    result, temporal = await ingest(db, user(role), authority_level=1)

    assert result["authority_level"] == 4 and result["authority_capped"] is True
    assert result["authority_requested"] == 1
    assert db.inserted("documents")[0]["authority_level"] == 4
    assert temporal.started[0]["authority_level"] == 4
    audit = db.inserted("audit_log")[0]["details"]
    assert audit["authority_requested"] == 1 and audit["authority_level"] == 4
    assert audit["authority_asserted_by"] == "u-1" and audit["uploader_role"] == role


@pytest.mark.parametrize("role", ["admin", "reliability"])
async def test_a_designated_role_may_assert_levels_one_to_three_and_is_recorded(role):
    db = FakeSupabase()
    result, temporal = await ingest(db, user(role, "boss-9"), authority_level=1)

    assert result["authority_level"] == 1 and result["authority_capped"] is False
    assert temporal.started[0]["authority_level"] == 1
    audit = db.inserted("audit_log")[0]["details"]
    assert audit["authority_asserted_by"] == "boss-9" and audit["uploader_role"] == role


async def test_levels_four_and_five_are_open_to_any_uploader():
    db = FakeSupabase()
    result, _ = await ingest(db, user("engineer"), authority_level=5)
    assert result["authority_level"] == 5 and result["authority_capped"] is False


async def test_an_unknown_asset_is_refused_before_anything_is_stored():
    db = FakeSupabase(assets=[{"asset_id": "HE-301"}])
    with pytest.raises(HTTPException) as exc:
        await ingest(db, user("admin"), asset_id="HE-999")
    assert exc.value.status_code == 422 and db.uploads == []

    result, _ = await ingest(db, user("admin"), asset_id="HE-301")
    assert result["status"] == "accepted"


@pytest.mark.parametrize("value", ["1990-01-01T00:00:00Z", "2999-01-01T00:00:00Z", "not a date"])
async def test_occurred_at_cannot_backdate_forward_date_or_be_garbage(value):
    db = FakeSupabase()
    with pytest.raises(HTTPException) as exc:
        await ingest(db, user("admin"), occurred_at=value)
    assert exc.value.status_code == 422 and db.uploads == []


def test_a_plausible_occurred_at_is_normalised_to_utc():
    assert docs.parse_occurred_at("2024-01-15T08:30:00Z", NOW) == "2024-01-15T08:30:00+00:00"
    assert docs.parse_occurred_at("2024-01-15T08:30:00", NOW) == "2024-01-15T08:30:00+00:00"
    assert docs.parse_occurred_at(None, NOW) is None


# ---------------------------------------------------------------------------
# L10 / L6: content type, signed URL, error text
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("sent,stored", [
    ("text/html", "application/octet-stream"),
    ("image/svg+xml", "application/octet-stream"),
    (None, "application/octet-stream"),
    ("application/pdf", "application/pdf"),
    ("Text/Plain; charset=utf-8", "text/plain"),
])
def test_only_known_document_types_are_stored(sent, stored):
    assert docs.normalise_mime(sent) == stored


async def test_the_stored_content_type_is_the_normalised_one():
    db = FakeSupabase()
    await ingest(db, user("admin"), file=upload("page.html", b"<script>1</script>", "text/html"))
    assert db.uploads[0][1] == {"content-type": "application/octet-stream"}
    assert db.inserted("documents")[0]["mime_type"] == "application/octet-stream"


async def test_signed_urls_force_a_download():
    url = f"https://x.supabase.co/storage/v1/object/authenticated/{docs.settings.SUPABASE_STORAGE_BUCKET}/procedure/DOC-1/a.pdf"
    db = FakeSupabase(documents=[{"document_id": "DOC-1", "vault_url": url, "file_name": "../a b.pdf"}])
    out = await docs.get_artifact_url("DOC-1", current_user=user("engineer"), supabase=db)
    assert out["signed_url"].endswith("&download=a_b.pdf")


async def test_a_storage_failure_returns_no_exception_text():
    class Boom(FakeSupabase):
        def __init__(self):
            super().__init__()
            self.storage = SimpleNamespace(from_=lambda b: SimpleNamespace(
                upload=lambda *a, **k: (_ for _ in ()).throw(RuntimeError("secret-bucket-detail"))))

    with pytest.raises(HTTPException) as exc:
        await ingest(Boom(), user("admin"))
    assert exc.value.status_code == 502 and "secret-bucket-detail" not in exc.value.detail


# ---------------------------------------------------------------------------
# H6: supersede
# ---------------------------------------------------------------------------

class FakeGraph:
    closed: list[str] = []

    def __init__(self, driver):
        pass

    async def get_blast_radius(self, document_id):
        return {"affected": [], "affected_count": 0}

    async def close_validity_windows_for_document(self, document_id, when):
        FakeGraph.closed.append(document_id)
        return 3


class OkIndex:
    async def update(self, **kwargs):
        return {}


class OkVectors:
    def __init__(self, *a):
        pass

    async def mark_superseded(self, *a):
        return None


@pytest.fixture
def supersede_env(monkeypatch):
    FakeGraph.closed = []
    monkeypatch.setattr(docs, "GraphService", FakeGraph)
    monkeypatch.setattr(docs, "VectorStoreService", OkVectors)


async def supersede(db, role, old="DOC-OLD", new="DOC-NEW"):
    response = Response()
    body = await docs.supersede_document(
        old, current_user=user(role), response=response, supabase=db, driver=None, es=OkIndex(),
        qdrant=None, settings_dep=docs.settings, new_document_id=new,
    )
    return body, response


def vault(old_authority=1, new_status="active", **extra):
    return FakeSupabase(documents=[
        {"document_id": "DOC-OLD", "status": "active", "authority_level": old_authority},
        {"document_id": "DOC-NEW", "status": new_status, "authority_level": 4, "ingested_by": "u-2"},
    ], **extra)


async def test_a_document_cannot_supersede_itself(supersede_env):
    with pytest.raises(HTTPException) as exc:
        await supersede(vault(), "admin", old="DOC-OLD", new="DOC-OLD")
    assert exc.value.status_code == 400


async def test_the_replacement_must_be_active(supersede_env):
    with pytest.raises(HTTPException) as exc:
        await supersede(vault(new_status="superseded"), "admin")
    assert exc.value.status_code == 409


async def test_an_engineer_cannot_supersede_a_regulatory_document(supersede_env):
    db = vault(old_authority=1)
    with pytest.raises(HTTPException) as exc:
        await supersede(db, "engineer")
    assert exc.value.status_code == 403
    assert db.writes == [] and FakeGraph.closed == []


async def test_a_gated_supersede_creates_the_moc_and_closes_nothing(supersede_env):
    db = vault(old_authority=2)
    body, response = await supersede(db, "reliability")

    assert response.status_code == 202 and body["status"] == "pending_moc_approval"
    moc = db.inserted("moc_items")[0]
    assert moc["status"] == "pending_approval" and moc["moc_id"] == body["moc_id"]
    assert db.updated("documents") == [] and FakeGraph.closed == []


async def test_repeating_a_pending_request_does_not_create_a_second_moc(supersede_env):
    moc_id = docs.supersede_moc_id("DOC-OLD", "DOC-NEW")
    db = vault(old_authority=1, moc_items=[{"moc_id": moc_id, "status": "pending_approval"}])
    body, response = await supersede(db, "admin")

    assert response.status_code == 202 and body["moc_id"] == moc_id
    assert db.inserted("moc_items") == [] and FakeGraph.closed == []


async def test_an_approved_moc_lets_the_supersession_apply(supersede_env):
    moc_id = docs.supersede_moc_id("DOC-OLD", "DOC-NEW")
    db = vault(old_authority=1, moc_items=[{"moc_id": moc_id, "status": "approved"}])
    body, _ = await supersede(db, "admin")

    assert body["status"] == "superseded" and body["moc_id"] == moc_id
    assert FakeGraph.closed == ["DOC-OLD"]
    assert {"status": "superseded"} in db.updated("documents")


async def test_a_rejected_moc_blocks_the_supersession(supersede_env):
    moc_id = docs.supersede_moc_id("DOC-OLD", "DOC-NEW")
    db = vault(old_authority=1, moc_items=[{"moc_id": moc_id, "status": "rejected"}])
    with pytest.raises(HTTPException) as exc:
        await supersede(db, "admin")
    assert exc.value.status_code == 409 and FakeGraph.closed == []


async def test_an_operational_document_still_supersedes_immediately(supersede_env):
    db = vault(old_authority=4)
    body, _ = await supersede(db, "engineer")
    assert body["status"] == "superseded" and FakeGraph.closed == ["DOC-OLD"]


# ---------------------------------------------------------------------------
# M11: PDF, XLSX and retry limits
# ---------------------------------------------------------------------------

def pdf_bytes(pages, size=(595, 842)):
    doc = fitz.open()
    for i in range(pages):
        doc.new_page(width=size[0], height=size[1]).insert_text((72, 72), f"page {i}")
    data = doc.tobytes()
    doc.close()
    return data


def test_a_pdf_with_too_many_pages_is_not_rasterised(monkeypatch):
    monkeypatch.setattr(ocr, "MAX_PDF_PAGES", 2)
    svc = ocr.OCRService()
    assert svc._to_images(pdf_bytes(3), "application/pdf") == []
    assert svc._extract_native_pdf(pdf_bytes(3))["blocks"] == []
    assert len(svc._to_images(pdf_bytes(2), "application/pdf")) == 2


def test_page_size_and_dpi_are_capped():
    a4 = ocr.capped_dpi(595, 842, 96)
    assert a4 == 96
    assert ocr.capped_dpi(595, 842, 600) == ocr.MAX_RASTER_DPI
    a0 = ocr.capped_dpi(2384, 3370, 150)
    assert a0 < 150 and (2384 / 72) * (3370 / 72) * a0 * a0 <= ocr.MAX_RASTER_PIXELS
    assert ocr.capped_dpi(14400, 14400, 150) is None
    assert ocr.capped_dpi(0, 100, 96) is None


def test_an_oversized_pdf_page_is_rejected_rather_than_rendered():
    assert ocr.OCRService()._to_images(pdf_bytes(1, size=(14400, 14400)), "application/pdf") == []


def test_the_pid_rasteriser_applies_the_same_ceiling():
    assert PIDService._first_image(pdf_bytes(1, size=(14400, 14400)), "application/pdf") is None
    png, mime = PIDService._first_image(pdf_bytes(1), "application/pdf")
    assert mime == "image/png" and png[:4] == b"\x89PNG"


def xlsx_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(rows):
        ws.append([f"row {i}", i])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_a_spreadsheet_over_the_row_cap_is_rejected_not_truncated(monkeypatch):
    monkeypatch.setattr(ocr, "MAX_SPREADSHEET_ROWS", 10)
    assert ocr.OCRService._extract_spreadsheet(xlsx_bytes(11)) == ""
    assert "row 9" in ocr.OCRService._extract_spreadsheet(xlsx_bytes(10))


def test_a_spreadsheet_that_inflates_past_the_cap_is_rejected(monkeypatch):
    monkeypatch.setattr(ocr, "MAX_SPREADSHEET_UNZIPPED_BYTES", 100)
    assert ocr.OCRService._extract_spreadsheet(xlsx_bytes(5)) == ""


def test_bad_file_failures_are_not_retried():
    assert "BadFile" in pipeline.DEFAULT_RETRY.non_retryable_error_types
    assert pipeline.DEFAULT_RETRY.maximum_attempts == 5  # transient failures keep their retries


async def test_an_integrity_failure_is_raised_as_non_retryable(monkeypatch):
    from temporalio.exceptions import ApplicationError

    class Storage:
        def from_(self, bucket):
            return SimpleNamespace(download=lambda path: b"tampered")

    db = FakeSupabase(documents=[{"document_id": "DOC-1", "sha256_hash": "0" * 64}])
    db.storage = Storage()
    monkeypatch.setattr(pipeline, "_get_supabase", lambda: db)
    with pytest.raises(ApplicationError) as exc:
        await pipeline.store_in_vault("DOC-1", "p/DOC-1/a.txt", "text/plain", "job-1")
    assert exc.value.non_retryable and exc.value.type == "BadFile"


# ---------------------------------------------------------------------------
# M13: only confirmed aliases link
# ---------------------------------------------------------------------------

async def test_the_pipeline_loads_confirmed_aliases_only(monkeypatch):
    db = FakeSupabase(documents=[{"document_id": "DOC-1", "document_type": "procedure"}])
    monkeypatch.setattr(pipeline, "_get_supabase", lambda: db)
    monkeypatch.setattr(pipeline, "_get_neo4j_driver", lambda: None)

    class Graph:
        def __init__(self, driver):
            pass

        async def merge_document_node(self, *a, **k):
            return None

    monkeypatch.setattr("api.services.graph.GraphService", Graph)
    await pipeline.link_to_graph("DOC-1", [], None, 4, "job-1")

    assert ("asset_alias_map", "confirmed", True) in db.eq_calls


def test_an_unconfirmed_alias_never_overwrites_an_existing_one():
    source = open(pipeline.__file__).read()
    assert 'on_conflict="alias", ignore_duplicates=True' in source


# ---------------------------------------------------------------------------
# L12: graph node ids are scoped to the drawing
# ---------------------------------------------------------------------------

def test_topology_node_ids_differ_across_drawings():
    a = GraphService.topology_node_id("DOC-A", "TOPO-EQ-001")
    b = GraphService.topology_node_id("DOC-B", "TOPO-EQ-001")
    assert a != b and a.startswith("DOC-A:") and a.endswith("TOPO-EQ-001")


# ---------------------------------------------------------------------------
# L9: wider PII patterns
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("leaked", ["1234-5678-9012", "+91 98765 43210", "+91-98765-43210", "(+91) 98765 43210",
                                    "91 9876543210", "98765 43210", "09876543210"])
def test_identifiers_in_other_common_formats_are_masked(leaked):
    out = PIIService().redact(f"Contact {leaked} for access.")
    assert leaked not in out["redacted_text"] and out["pii_found"]


def test_names_first_mentioned_late_in_a_long_document_are_masked():
    text = ("Routine inspection of pump EQ-101. " * 400) + "Closed out by Mr. Rajesh Kumar. Signed by: Anita Rao."
    out = PIIService().redact(text, person_names=[])
    assert "Rajesh Kumar" not in out["redacted_text"] and "Anita Rao" not in out["redacted_text"]
    assert out["counts"]["PERSON"] == 2


def test_operational_content_is_still_untouched():
    text = "Pump EQ-101 seal P/N MS-4471-B per OISD-117. Engineer: signoff required. Valve XV-203, loop 1234-56."
    assert PIIService().redact(text)["redacted_text"] == text
